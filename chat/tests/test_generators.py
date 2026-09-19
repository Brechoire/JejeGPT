import io

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from chat.exports.docx_gen import build_docx
from chat.exports.pdf_gen import build_pdf
from chat.exports.pptx_gen import build_pptx
from chat.exports.xlsx_gen import build_xlsx

SAMPLE = """# Fréquentation

Paragraphe d'intro avec accents été.

- Atelier lecture
- Accueil de classes

| Mois | Entrées |
|------|---------|
| Janvier | 120 |
| Février | 150 |

```python
print("hello")
```
"""


def test_docx_genere_reouvrable():
    payload = build_docx("Rapport", SAMPLE)
    assert payload[:2] == b"PK"
    doc = Document(io.BytesIO(payload))
    texts = [p.text for p in doc.paragraphs]
    assert any("Rapport" in t for t in texts)
    assert any("Fréquentation" in t for t in texts)
    assert len(doc.tables) == 1


def test_xlsx_genere_formules_et_sanitization():
    payload = build_xlsx("Stats", SAMPLE + "\n| =1+1 | x |\n|---|---|\n| =cmd | y |\n")
    assert payload[:2] == b"PK"
    wb = load_workbook(io.BytesIO(payload))
    assert "Contenu" in wb.sheetnames
    assert "Tableau 1" in wb.sheetnames
    ws = wb["Contenu"]
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any(v.startswith("'=") for v in values)
    assert any("SUM(" in v or "AVERAGE(" in v for v in values)


def test_pptx_genere_reouvrable():
    payload = build_pptx("Présentation", SAMPLE)
    assert payload[:2] == b"PK"
    prs = Presentation(io.BytesIO(payload))
    assert len(prs.slides) >= 3


def test_pdf_genere_et_echappe_html():
    payload = build_pdf("Rapport <test>", "Texte avec <b>balise</b> & esperluette.")
    assert payload[:5] == b"%PDF-"
    assert len(payload) > 1500


def test_generateurs_refusent_vide():
    for builder in (build_docx, build_xlsx, build_pptx, build_pdf):
        with pytest.raises(ValueError):
            builder("   ", "   ")

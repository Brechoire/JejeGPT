"""Rendu Markdown dans les exports : plus de syntaxe brute (#, **, `, [](), >, ---)."""
import io

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from chat.exports.document import CodeBlock, Heading, OrderedList, parse_markdown, strip_inline
from chat.exports.docx_gen import build_docx
from chat.exports.pdf_gen import _md, build_pdf
from chat.exports.pptx_gen import build_pptx
from chat.exports.xlsx_gen import build_xlsx

SAMPLE = """# Titre principal

Texte avec **gras**, *italique*, `code`, ~~barré~~ et [lien](https://example.com).

## Sous-titre

1. Premier **gras**
2. Deuxième `code`

- [ ] tâche à faire
- item *italique*

> citation importante

---

| Nom | Valeur |
|-----|--------|
| **clé** | `123` |

```python
print('verbatim')
```
"""

RESIDUS = ["**", "*italique*", "`code`", "[](", "# Titre", "> citation", "- [ ]"]


def test_parse_blocks_riches():
    blocks = parse_markdown(SAMPLE)
    kinds = [type(b).__name__ for b in blocks]
    assert "Heading" in kinds
    assert "OrderedList" in kinds
    assert "BulletList" in kinds
    assert "Quote" in kinds
    assert "HorizontalRule" in kinds
    assert "Table" in kinds
    assert "CodeBlock" in kinds
    ordered = next(b for b in blocks if isinstance(b, OrderedList))
    assert ordered.items == ["Premier **gras**", "Deuxième `code`"]
    code = next(b for b in blocks if isinstance(b, CodeBlock))
    assert code.code == "print('verbatim')"
    assert code.language == "python"


def test_parse_titres_niveaux_et_sans_espace():
    blocks = parse_markdown("# A\n\n#### Profond\n\n#SansEspace\n")
    assert [(b.level, b.text) for b in blocks if isinstance(b, Heading)] == [
        (1, "A"),
        (4, "Profond"),
        (1, "SansEspace"),
    ]


def test_strip_inline():
    assert strip_inline("**gras** et *italique*") == "gras et italique"
    assert strip_inline("`code` et ~~barré~~") == "code et barré"
    assert strip_inline("[lien](https://example.com)") == "lien"
    assert strip_inline("Voir ![photo](https://img/x.png) ici") == "Voir photo ici"
    assert strip_inline("<b>html</b> & texte") == "<b>html</b> & texte"


def test_md_reportlab_sans_syntaxe():
    out = _md("**gras**, *italique*, `code` et [lien](https://example.com).")
    assert "**" not in out and "[](" not in out and "`" not in out
    assert "<b>gras</b>" in out
    assert "<i>italique</i>" in out
    assert '<link href="https://example.com">lien</link>' in out


def test_docx_sans_syntaxe_markdown():
    doc = Document(io.BytesIO(build_docx("Titre", SAMPLE)))
    texts = [p.text for p in doc.paragraphs]
    assert any("Titre principal" in t for t in texts)
    assert any("Premier gras" in t for t in texts)
    assert any("citation importante" in t for t in texts)
    blob = "\n".join(t for t in texts if "verbatim" not in t)
    for marqueur in RESIDUS:
        assert marqueur not in blob
    assert any("print('verbatim')" in t for t in texts)


def test_xlsx_sans_syntaxe_et_sanitization():
    payload = build_xlsx("Titre", SAMPLE + "\n| =1+1 | x |\n|---|---|\n| =cmd | y |\n")
    ws = load_workbook(io.BytesIO(payload))["Contenu"]
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("Premier gras" in v for v in values)
    assert any(v.startswith("'=") for v in values)
    blob = "\n".join(v for v in values if "verbatim" not in v)
    for marqueur in RESIDUS:
        assert marqueur not in blob


def test_pptx_sans_syntaxe_markdown():
    prs = Presentation(io.BytesIO(build_pptx("Titre", SAMPLE)))
    texts = [getattr(sh, "text", "") for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert any("Titre principal" in t for t in texts)
    blob = "\n".join(t for t in texts if "verbatim" not in t)
    for marqueur in RESIDUS:
        assert marqueur not in blob


def test_pdf_construit_markdown_riche():
    payload = build_pdf("Rapport", SAMPLE)
    assert payload[:5] == b"%PDF-"
    assert len(payload) > 2000

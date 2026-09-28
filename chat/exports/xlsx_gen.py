"""Générateur XLSX multi-feuilles (contenu, tableaux, totaux avec formules natives)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .document import (
    Block,
    BulletList,
    CodeBlock,
    Heading,
    HorizontalRule,
    OrderedList,
    Paragraph,
    Quote,
    Table,
    parse_markdown,
    sanitize_cell,
    strip_inline,
)

_TITLE_FONT = Font(bold=True, size=14)
_HEADING_FONT = Font(bold=True, size=12)
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _is_number(text: str) -> bool:
    try:
        float(text.replace(" ", "").replace(",", "."))
        return True
    except ValueError:
        return False


def _write_table(ws, table: Table, start_row: int = 1) -> int:
    row = start_row
    for ci, header in enumerate(table.headers, start=1):
        cell = ws.cell(row=row, column=ci, value=sanitize_cell(strip_inline(header)))
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
    row += 1
    numeric_cols = [all(_is_number(r[ci]) for r in table.rows if ci < len(r)) for ci in range(len(table.headers))]
    for values in table.rows:
        for ci in range(len(table.headers)):
            raw = strip_inline(values[ci] if ci < len(values) else "")
            text = sanitize_cell(raw)
            if _is_number(raw):
                try:
                    ws.cell(row=row, column=ci + 1, value=float(raw.replace(" ", "").replace(",", ".")))
                except ValueError:
                    ws.cell(row=row, column=ci + 1, value=text)
            else:
                ws.cell(row=row, column=ci + 1, value=text)
        row += 1
    if table.rows and any(numeric_cols):
        ws.cell(row=row, column=1, value="Total").font = _HEADER_FONT
        ws.cell(row=row + 1, column=1, value="Moyenne").font = _HEADER_FONT
        for ci, numeric in enumerate(numeric_cols, start=1):
            if numeric:
                col = get_column_letter(ci)
                first = start_row + 1
                last = row - 1
                ws.cell(row=row, column=ci, value=f"=SUM({col}{first}:{col}{last})")
                ws.cell(row=row + 1, column=ci, value=f"=AVERAGE({col}{first}:{col}{last})")
        row += 2
    for ci in range(1, len(table.headers) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 22
    return row + 1


def build_xlsx(title: str, markdown_text: str) -> bytes:
    blocks: list[Block] = parse_markdown(markdown_text)
    if not title.strip() and not blocks:
        raise ValueError("Contenu vide.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Contenu"
    ws["A1"] = strip_inline(title.strip()) or "Export JejeGPT"
    ws["A1"].font = _TITLE_FONT
    row = 3
    table_count = 0
    for block in blocks:
        if isinstance(block, Heading):
            ws.cell(row=row, column=1, value=strip_inline(block.text)).font = _HEADING_FONT
            row += 1
        elif isinstance(block, Paragraph):
            ws.cell(row=row, column=1, value=strip_inline(block.text)).alignment = _WRAP
            row += 1
        elif isinstance(block, Quote):
            ws.cell(row=row, column=1, value="« " + strip_inline(block.text) + " »").alignment = _WRAP
            row += 1
        elif isinstance(block, HorizontalRule):
            ws.cell(row=row, column=1, value="———").alignment = _WRAP
            row += 1
        elif isinstance(block, BulletList):
            for item in block.items:
                ws.cell(row=row, column=1, value="• " + sanitize_cell(strip_inline(item))).alignment = _WRAP
                row += 1
        elif isinstance(block, OrderedList):
            for n, item in enumerate(block.items, start=1):
                ws.cell(row=row, column=1, value=f"{n}. " + sanitize_cell(strip_inline(item))).alignment = _WRAP
                row += 1
        elif isinstance(block, Table):
            table_count += 1
            end = _write_table(ws, block, start_row=row)
            detail = wb.create_sheet(title=f"Tableau {table_count}")
            _write_table(detail, block)
            row = end
        elif isinstance(block, CodeBlock):
            ws.cell(row=row, column=1, value=sanitize_cell(block.code)).alignment = _WRAP
            row += 1
    ws.column_dimensions["A"].width = 80
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

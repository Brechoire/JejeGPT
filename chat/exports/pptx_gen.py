"""Générateur PPTX natif (title slide, sections, listes, tableaux)."""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.util import Inches

from .document import (
    Block,
    BulletList,
    CodeBlock,
    Heading,
    HorizontalRule,
    OrderedList,
    Paragraph,
    Quote,
    Table,
    parse_markdown,
    strip_inline,
)


def _add_table(slide, table: Table) -> None:
    rows = [table.headers, *table.rows]
    cols = max(len(r) for r in rows)
    left, top, width = Inches(0.5), Inches(1.4), Inches(9)
    shape = slide.shapes.add_table(len(rows), cols, left, top, width, Inches(0.4))
    for ri, row in enumerate(rows):
        for ci in range(cols):
            shape.table.cell(ri, ci).text = strip_inline(row[ci]) if ci < len(row) else ""


def _add_bullets(slide, items: list[str], ordered: bool = False) -> None:
    tx_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(9), Inches(4))
    frame = tx_box.text_frame
    frame.word_wrap = True
    for n, item in enumerate(items):
        para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        prefix = f"{n + 1}. " if ordered else "• "
        para.text = prefix + strip_inline(item)
        para.level = 0


def _add_text(slide, text: str) -> None:
    tx_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(9), Inches(4))
    frame = tx_box.text_frame
    frame.word_wrap = True
    frame.paragraphs[0].text = text


def build_pptx(title: str, markdown_text: str) -> bytes:
    blocks: list[Block] = parse_markdown(markdown_text)
    if not title.strip() and not blocks:
        raise ValueError("Contenu vide.")
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = strip_inline(title.strip()) or "Export JejeGPT"
    title_slide.placeholders[1].text = "Généré par JejeGPT — usage local"
    current_title = strip_inline(title.strip()) or "Contenu"
    pending_text: list[str] = []

    def flush_text() -> None:
        if pending_text:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = current_title[:80]
            _add_text(slide, "\n\n".join(pending_text))
            pending_text.clear()

    for block in blocks:
        if isinstance(block, Heading):
            flush_text()
            current_title = strip_inline(block.text)
        elif isinstance(block, Paragraph):
            pending_text.append(strip_inline(block.text))
        elif isinstance(block, Quote):
            pending_text.append("« " + strip_inline(block.text) + " »")
        elif isinstance(block, HorizontalRule):
            pending_text.append("———")
        elif isinstance(block, BulletList):
            flush_text()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = current_title[:80]
            _add_bullets(slide, block.items)
        elif isinstance(block, OrderedList):
            flush_text()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = current_title[:80]
            _add_bullets(slide, block.items, ordered=True)
        elif isinstance(block, Table):
            flush_text()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = current_title[:80]
            _add_table(slide, block)
        elif isinstance(block, CodeBlock):
            pending_text.append(f"[{block.language}]\n{block.code}")
    flush_text()
    if len(prs.slides) == 1:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = current_title[:80]
        _add_text(slide, "(Aucun contenu structuré)")
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()

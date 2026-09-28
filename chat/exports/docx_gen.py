"""Générateur DOCX structuré (titres, paragraphes, listes, tableaux, code, pied de page)."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .document import (
    Block,
    BulletList,
    CodeBlock,
    Heading,
    HorizontalRule,
    OrderedList,
    Paragraph,
    Quote,
    Table,
    inline_segments,
    parse_markdown,
    strip_inline,
)


def _add_hyperlink(paragraph, url: str, text: str):
    """Ajoute un lien cliquable (python-docx n'expose pas d'API native)."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_runs(paragraph, text: str) -> None:
    for seg in inline_segments(text or ""):
        if seg.link and not seg.code:
            _add_hyperlink(paragraph, seg.link, seg.text)
            continue
        run = paragraph.add_run(seg.text)
        run.bold = seg.bold or None
        run.italic = seg.italic or None
        run.font.strike = seg.strike or None
        if seg.code:
            run.font.size = Pt(9)
            run.font.name = "Consolas"


def _add_rich_paragraph(doc, text: str):
    para = doc.add_paragraph()
    _add_runs(para, text)
    return para


def _set_cell_text(cell, text: str) -> None:
    cell.text = ""
    _add_runs(cell.paragraphs[0], text)


def build_docx(title: str, markdown_text: str) -> bytes:
    blocks: list[Block] = parse_markdown(markdown_text)
    if not title.strip() and not blocks:
        raise ValueError("Contenu vide.")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    doc.add_heading(strip_inline(title.strip()) or "Export JejeGPT", level=0)
    for block in blocks:
        if isinstance(block, Heading):
            doc.add_heading(strip_inline(block.text), level=min(block.level, 3))
        elif isinstance(block, Paragraph):
            _add_rich_paragraph(doc, block.text)
        elif isinstance(block, Quote):
            para = doc.add_paragraph(style="Intense Quote")
            _add_runs(para, block.text)
        elif isinstance(block, HorizontalRule):
            doc.add_paragraph("———")
        elif isinstance(block, BulletList):
            for item in block.items:
                para = doc.add_paragraph(style="List Bullet")
                _add_runs(para, item)
        elif isinstance(block, OrderedList):
            for item in block.items:
                para = doc.add_paragraph(style="List Number")
                _add_runs(para, item)
        elif isinstance(block, Table):
            rows = [block.headers, *block.rows]
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Light Grid Accent 1"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    cell = table.cell(ri, ci)
                    _set_cell_text(cell, row[ci] if ci < len(row) else "")
                    if ri == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
        elif isinstance(block, CodeBlock):
            para = doc.add_paragraph()
            run = para.add_run(block.code)
            run.font.size = Pt(9)
            run.font.name = "Consolas"
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = "Généré par JejeGPT — usage local"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
